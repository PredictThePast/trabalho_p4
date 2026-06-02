from openpyxl import load_workbook

def verificar_presencas(caminho):
    wb = load_workbook(caminho)
    ws = wb["Sheet1"]

    ws["D1"] = "Percentagem"
    ws["E1"] = "Situação"

    for i in range(2, ws.max_row + 1):
        presencas = ws[f"B{i}"].value
        total = ws[f"C{i}"].value

        percentagem = (presencas / total) * 100

        ws[f"D{i}"] = round(percentagem, 1)

        if percentagem >= 75:
            ws[f"E{i}"] = "Admitido"
        else:
            ws[f"E{i}"] = "Excluído"

    wb.save(caminho)
    print("Presenças verificadas.")

verificar_presencas("alunos.xlsx")