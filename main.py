from openpyxl import load_workbook

def avaliar_credito(caminho):
    wb = load_workbook(caminho)
    ws = wb["Sheet1"]

    ws["D1"] = "Rácio Dívida"
    ws["E1"] = "Decisão"

    for i in range(2, ws.max_row + 1):
        rendimento = ws[f"B{i}"].value
        divida = ws[f"C{i}"].value

        if rendimento and rendimento > 0:
            racio = divida / rendimento
        else:
            racio = 1

        ws[f"D{i}"] = round(racio, 2)

        if racio < 0.3:
            ws[f"E{i}"] = "Aprovado"
        elif racio < 0.6:
            ws[f"E{i}"] = "Análise"
        else:
            ws[f"E{i}"] = "Recusado"

    wb.save(caminho)
    print("Avaliação de crédito concluída.")

avaliar_credito("credito.xlsx")